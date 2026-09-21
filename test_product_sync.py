"""Run with python -m unittest test_product_sync -v (no network required)."""
import io
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook
import product_sync as sync
from build_codes import DICT_SHEETS, read_workbook
import build


def workbook(description='Test product', missing_sheet=False):
    wb = Workbook()
    wb.remove(wb.active)
    for name in DICT_SHEETS[:-1] if missing_sheet else DICT_SHEETS:
        ws = wb.create_sheet(name)
        ws.append(['CODE', 'DESCRIPTION', 'MODEL', 'COST', 'LABOR', 'MISC.'])
        ws.append(['EXIT DEVICES', None, None, None, None, 'Reference link'])
        ws.append(['NEW1', description, 'Example model'])
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    return stream.getvalue()


class ProductSyncTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.patches = [patch.object(sync, 'CACHE', Path(self.temp.name) / 'cache.json'),
                        patch.object(sync, '_data', None),
                        patch.object(sync, '_last_attempt', None),
                        patch.object(sync, '_synced_at', None),
                        patch.object(sync, '_error', None)]
        for p in self.patches:
            p.start()
            self.addCleanup(p.stop)

    def test_live_update_and_saved_cache(self):
        with patch.object(sync, '_download', return_value=workbook()):
            first, status = sync.get_products()
        self.assertEqual(status['state'], 'live')
        self.assertEqual(len(first), 14)
        self.assertEqual(first['Brivo'][0]['category'], 'EXIT DEVICES')
        with patch.object(sync, '_download', return_value=workbook('Updated wording')):
            updated, after = sync.get_products(force=True)
        self.assertNotEqual(status['version'], after['version'])
        self.assertEqual(updated['Brivo'][0]['description'], 'Updated wording')
        sync._data = None  # simulate restarting offline
        sync._synced_at = None
        with patch.object(sync, '_download', side_effect=OSError('Offline')):
            restored, offline = sync.get_products(force=True)
        self.assertEqual(restored, updated)
        self.assertEqual(offline['state'], 'cached')
        self.assertEqual(offline['last_synced'], after['last_synced'])

    def test_bad_export_keeps_cache_and_recovery_works(self):
        with patch.object(sync, '_download', return_value=workbook()):
            first, _ = sync.get_products()
        original_bytes = sync.CACHE.read_bytes()
        for invalid in (b'<html>Sign in</html>', workbook(missing_sheet=True)):
            with patch.object(sync, '_download', return_value=invalid):
                kept, status = sync.get_products(force=True)
            self.assertEqual(kept, first)
            self.assertEqual(status['state'], 'cached')
            self.assertEqual(sync.CACHE.read_bytes(), original_bytes)
        with patch.object(sync, '_download', return_value=workbook('Recovered')):
            _, status = sync.get_products(force=True)
        self.assertEqual(status['state'], 'live')
        self.assertIsNone(status['error'])

    def test_offline_first_run_and_rate_limit(self):
        with patch.object(sync, '_download', side_effect=TimeoutError('Timed out')) as download:
            data, status = sync.get_products()
            sync.get_products()
        self.assertEqual(download.call_count, 1)
        self.assertTrue(data['Brivo'])
        self.assertIsNone(status['last_synced'])
        self.assertEqual(status['state'], 'cached')

    def test_missing_description_is_rejected(self):
        with self.assertRaisesRegex(ValueError, 'Missing description'):
            read_workbook(io.BytesIO(workbook('')), strict=True)

    def test_builder_uses_live_products_and_preserves_saved_wording(self):
        with patch.object(sync, '_download', return_value=workbook('Fresh product')):
            sync.get_products()
        _, index = build.load_codes()
        self.assertEqual(build.resolve_line(index, {'code': 'NEW1', 'sheet': 'Brivo'})['text'],
                         'Fresh product.')
        self.assertEqual(build.resolve_line(index, {'code': 'REMOVED', 'text': 'Saved wording'})['text'],
                         'Saved wording.')
        self.assertEqual(build.resolve_codes(index, [{'code': 'OLD', 'text': 'Saved note'}], 'note'),
                         ['Saved note'])

    def test_proposal_generation_uses_new_products(self):
        import app
        from docx import Document
        with patch.object(sync, '_download', return_value=workbook('New live product')):
            response = app.api_codes(refresh=True)
        self.assertEqual(response['sync']['state'], 'live')
        job = {'proposal': {'for': 'Sync verification', 'tariff_notes': [], 'intro': 'Test intro'},
               'gates': [{'title': 'Test location', 'lines': [
                   {'code': 'NEW1', 'sheet': 'Brivo', 'qty': 1},
                   {'code': 'REMOVED', 'text': 'Preserved quote wording', 'qty': 1}]}],
               'notes': [{'code': 'OLD', 'text': 'Preserved note'}]}
        with patch.object(app, 'JOBS', self.temp.name):
            result = app.api_build(job)
        self.assertIsInstance(result, dict, getattr(result, 'body', None))
        self.assertTrue(result['ok'])
        document = Document(Path(self.temp.name) / result['slug'] / result['docx'])
        xml = document.element.xml
        for expected in ('New live product', 'Preserved quote wording', 'Preserved note'):
            self.assertIn(expected, xml)


if __name__ == '__main__':
    unittest.main()
