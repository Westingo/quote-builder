"""Fixed PM proposal wording transcribed from the highlighted Metro reference."""
FREQUENCIES = {'Monthly': 12, 'Quarterly': 4, 'Semiannual': 2, 'Annual': 1}
SCOPE = (
    'Includes lubrication, adjustments and inspection for wear',
    'Inspection & testing of all gate related safety devices',
    'If need for repairs are found, they will be quoted first',
    'Please have your pest control provider maintain the area surrounding the operators to prevent interference from insects/ small animals.',
)
NOTES = (
    'If UL 325-required safety warning signs are unreadable or missing, customer authorizes technician to replace signs at a cost of $20/per sign. Pricing subject to change without warning as contract auto-renews.',
    'Replace batteries for MATDCBB3, CSL24V or CSW24V operator while on site for PM if battery test falls below 10% of rated value at a cost of $33/each.',
    'Work will be billed after completion - do not send payment until then.',
    'This service does not guarantee that there will not be other repairs needed - It just keeps your system well-maintained.',
    'Hourly rates for other service call outs - Minimum service call (based on ½ travel, ½ hour on the job) $180.00. After the first hour, we bill in quarter hour increments based on the $180/hour rate. Rates subject to change. After hours rates are $281/hour with a minimum of 2 hours billed.',
    'Service contract automatically renews unless cancelled by customer prior to next service.',
    'Taxes, if applicable, not included in above totals and will be added to final billing.',
    'If a contract in excess of two pages is required, an additional $350 fee will apply',
    'Preventive maintenance services are set with flexible timing and the total represents a discounted billing rate as these services are designed to be completed while in the area for other work (without advance scheduling).',
    'Service accounts must be in good standings to continue our Preventative Maintenance agreement. We reserve the right to cancel any agreement that has an extended past due balance',
    'Due to current supply chain limitations, shipping and manufacturing times are unknown and will need to be verified upon approval of project.',
)


def validate(value):
    if not isinstance(value, dict):
        raise ValueError('PM quote details must be an object.')
    result = {}
    for key, default in [('frequency', 'Quarterly'), ('equipment', ''), ('price', '')]:
        item = value.get(key, default)
        if not isinstance(item, str) or len(item) > 10000:
            raise ValueError(f'PM {key} must be text of at most 10,000 characters.')
        result[key] = item.strip()
    if result['frequency'] not in FREQUENCIES:
        raise ValueError('Choose a valid PM frequency.')
    return result


def render(body, value):
    import proposal as p
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = validate(value)
    p._band(body, 'PREVENTATIVE MAINTENANCE PROPOSAL')
    heading = p._para(body, before=8, after=8)
    heading.paragraph_format.keep_with_next = True
    p._run(heading, f"Recommended {d['frequency']} Preventative Maintenance Service for the following:", bold=True, underline=True, size=10)
    for equipment in d['equipment'].splitlines():
        if equipment.strip():
            para = p._para(body, after=2)
            p._run(para, '• ' + equipment.strip(), size=10)
    for text in SCOPE:
        para = p._para(body, after=1)
        para.paragraph_format.left_indent = Inches(.5)
        p._run(para, '• ' + text, size=10)
    visits = FREQUENCIES[d['frequency']]
    table = body.add_table(rows=1, cols=2)
    p.U.clear_table_borders(table)
    p.U.set_col_widths(table, [p.CONTENT_W, p.AMOUNT_W])
    p.U.cant_split_row(table.rows[0])
    for cell in table.rows[0].cells:
        p.U.no_space(cell.paragraphs[0], after=0)
    label = table.cell(0, 0).paragraphs[0]
    label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p._run(label, f"Cost Per Year ({visits} {'Visit' if visits == 1 else 'Visits'}):", bold=True, underline=True, size=10)
    amount = p._amount_cell(table.cell(0, 1), d['price'])
    amount.paragraph_format.left_indent = Inches(.08)
    heading = p._para(body, before=8, after=6)
    heading.paragraph_format.keep_with_next = True
    p._run(heading, 'Notes:', bold=True, underline=True, size=10)
    for text in NOTES:
        para = p._para(body, after=1)
        if 'do not send payment until then.' in text:
            before, after = text.split('do not send payment until then.')
            p._run(para, '- ' + before, size=9.5)
            p._run(para, 'do not send payment until then.', bold=True, size=9.5)
        else:
            p._run(para, '- ' + text, size=9.5)
