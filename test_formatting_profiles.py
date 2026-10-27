"""John's experiments must not affect a standard quote in the same process."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
import build
import proposal
import quote_docx
import quote_transfer
import john_formatting


class ProfileTests(unittest.TestCase):
    def test_isolation_and_portable_selection(self):
        data, index = build.load_codes()
        standard = {'proposal': {'for': 'Profile verification'}, 'gates': []}
        john = dict(standard, formatting_profile='john')
        with tempfile.TemporaryDirectory() as folder:
            first, test, last = [Path(folder) / name for name in ('first.docx', 'john.docx', 'last.docx')]
            proposal.build_proposal(build.build_doc(standard, data, index), str(first))
            with patch.object(john_formatting, 'apply', side_effect=lambda doc: doc.add_paragraph('JOHN ONLY')) as hook:
                proposal.build_proposal(build.build_doc(john, data, index), str(test))
                proposal.build_proposal(build.build_doc(standard, data, index), str(last))
                self.assertEqual(hook.call_count, 1)
            self.assertIn('JOHN ONLY', Document(test).element.xml)
            self.assertEqual(Document(first).element.xml, Document(last).element.xml)
            quote_docx.embed(test, john)
            self.assertEqual(quote_docx.extract(test.read_bytes())['formatting_profile'], 'john')
        exported = quote_transfer.export_quote(john)
        self.assertEqual(json.loads(exported)['version'], 3)
        self.assertEqual(quote_transfer.import_quote(exported), john)
        self.assertEqual(json.loads(quote_transfer.export_quote(standard))['version'], 1)
        with self.assertRaises(ValueError):
            quote_transfer.export_quote(dict(standard, formatting_profile='unknown'))


if __name__ == '__main__':
    unittest.main()
