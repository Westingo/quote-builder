import unittest
import json
import build
import pm_quotes
import proposal
import quote_transfer
from pathlib import Path
from tempfile import TemporaryDirectory
from docx import Document

class PMQuoteTests(unittest.TestCase):
    def job(self, frequency='Quarterly'):
        return {'quote_type':'pm','proposal':{'for':'PM test'}, 'pm_quote':{'frequency':frequency,'equipment':'2) Slide Gates with 2 LiftMaster Model SL3000 AC Slide Gate Operators','price':'1,440.00'}}

    def test_frequency_and_fixed_wording_in_document(self):
        for frequency, visits in pm_quotes.FREQUENCIES.items():
            job=self.job(frequency)
            doc=build.build_doc(job, {}, {})
            with TemporaryDirectory() as folder:
                path=Path(folder)/'pm.docx'
                proposal.build_proposal(doc,path)
                output=Document(path)
                text='\n'.join([p.text for p in output.paragraphs]+[c.text for t in output.tables for r in t.rows for c in r.cells])
                self.assertIn(f'Recommended {frequency}',text)
                self.assertIn(f"Cost Per Year ({visits} {'Visit' if visits == 1 else 'Visits'}):",text)
                self.assertIn('$ 1,440.00',text)
                for fixed in (*pm_quotes.SCOPE,*pm_quotes.NOTES):
                    self.assertIn(fixed,text)
                self.assertNotIn('WE PROPOSE TO FURNISH',text)
                self.assertNotIn('Install:',text)

    def test_round_trip_and_version(self):
        job=self.job('Semiannual')
        encoded=quote_transfer.export_quote(job)
        self.assertEqual(json.loads(encoded)['version'],4)
        self.assertEqual(quote_transfer.import_quote(encoded),job)

    def test_invalid_frequency_and_types_rejected(self):
        for field,value in [('frequency','Weekly'),('price',[]),('equipment',{})]:
            job=self.job();job['pm_quote'][field]=value
            with self.assertRaises(ValueError):quote_transfer.export_quote(job)

    def test_blank_draft_renders(self):
        self.assertEqual(build.build_doc({'quote_type':'pm'}, {}, {})['pm_quote']['price'],'')

if __name__=='__main__':unittest.main()
