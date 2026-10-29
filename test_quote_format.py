"""Formatting survives transfers and reaches actual Word paragraph/cell settings."""
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.shared import Inches
from docx.enum.table import WD_ALIGN_VERTICAL

import build
import proposal
import quote_transfer
import quote_format


def fixture():
    return {"proposal": {"for": "Formatting check"}, "gates": [
        {"title": "Aligned items", "format": {"font_size": 9, "space_after": 5}, "lines": [
            {"qty": 2, "text": "First description line\nSecond description line\nThird description line", "amount": "100", "format": {"amount_align": "top", "bold": True, "italic": True}},
            {"qty": 1, "text": "Text starts at the quantity column and wraps to that same position. " * 3, "amount": "200", "format": {"start": "quantity", "wrap": "start", "amount_align": "center", "underline": True}},
            {"text": "Full-width note with no quantity marker. " * 3, "format": {"start": "left", "font_size": 10}}]},
        {"title": "New-page section", "format": {"page_break": True, "keep_together": True}, "lines": [{"text": "Section kept together", "qty": 1}]}],
        "options": [{"title": "Detailed option", "format": {"font_size": 9}, "lines": [{"amount_note": "Amount at top", "amount": "300", "format": {"amount_align": "top"}}]}]}


class FormatTests(unittest.TestCase):
    def test_transfer_and_docx_formatting(self):
        job = fixture()
        self.assertEqual(quote_transfer.import_quote(quote_transfer.export_quote(job)), job)
        data, index = build.load_codes()
        normalized = build.build_doc(job, data, index)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "test.docx"
            proposal.build_proposal(normalized, str(path))
            doc = Document(path)
        gate = doc.tables[1]
        first = gate.cell(1, 0).paragraphs[0]
        self.assertTrue(all(r.bold for r in first.runs))
        self.assertTrue(all(r.italic for r in first.runs))
        self.assertEqual(first.runs[-1].font.size.pt, 9)
        self.assertIn("\nSecond", first.text)
        self.assertEqual(first.paragraph_format.space_after.pt, 5)
        self.assertEqual(gate.cell(1, 1).vertical_alignment, WD_ALIGN_VERTICAL.TOP)
        second = gate.cell(2, 0).paragraphs[0]
        self.assertTrue(all(r.underline for r in second.runs))
        self.assertAlmostEqual(second.paragraph_format.left_indent.inches, 0.62, places=3)
        self.assertNotIn("1)", second.text)
        self.assertEqual(gate.cell(2, 1).vertical_alignment, WD_ALIGN_VERTICAL.CENTER)
        third = gate.cell(3, 0).paragraphs[0]
        self.assertEqual(third.paragraph_format.left_indent, 0)
        self.assertTrue(doc._element.xpath('.//w:br[@w:type="page"]'))
        self.assertTrue(doc.tables[2].cell(0, 0).paragraphs[0].paragraph_format.keep_with_next)
        self.assertEqual(doc.tables[3].cell(1, 1).vertical_alignment, WD_ALIGN_VERTICAL.TOP)

    def test_supply_group_spacing(self):
        lines = [
            {"qty": 1, "text": "Installed opener", "label": "Install:"},
            {"qty": 10, "text": "Keychain remotes", "label": "Supply:"},
            {"qty": 1, "text": "Manuals", "label": ""},
            {"qty": 1, "text": "Custom spacing", "label": "Supply:",
             "format": {"space_before": 6}},
        ]
        doc = Document()
        proposal._gate_block(doc, {"title": "Gate", "lines": lines})
        proposal._detail_option(doc, {"title": "Option", "lines": lines})
        for table in doc.tables:
            for column in (0, 1):
                paras = [table.cell(i, column).paragraphs[0] for i in range(1, 5)]
                self.assertEqual([p.paragraph_format.space_before.pt for p in paras],
                                 [0, 12, 0, 6])

    def test_cmi_badge_and_transfer(self):
        data, index = build.load_codes()
        for enabled in (True, False):
            job = {"proposal": {"for": "Property name", "address": "Property address", "cmi": enabled}}
            self.assertEqual(quote_transfer.import_quote(quote_transfer.export_quote(job)), job)
            normalized = build.build_doc(job, data, index)
            self.assertEqual(normalized["header"]["for"], "Property name")
            self.assertEqual(normalized["header"]["address"], "Property address")
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / "cmi.docx"
                proposal.build_proposal(normalized, str(path))
                doc = Document(path)
            for header in (doc.sections[0].header, doc.sections[0].first_page_header):
                self.assertEqual('CMIBadge' in header._element.xml, enabled)

    def test_summary_text_formatting(self):
        job = {"proposal": {"for": "Summary test"}, "summary_text": "Project details\nSecond line",
               "summary_text_format": {"bold": True, "italic": True, "underline": True,
                   "align": "center", "font_size": 12, "space_before": 10, "space_after": 8,
                   "line_spacing": 1.2, "keep_next": True, "start": "quantity", "wrap": "description"}}
        self.assertEqual(quote_transfer.import_quote(quote_transfer.export_quote(job)), job)
        data, index = build.load_codes()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "summary.docx"
            proposal.build_proposal(build.build_doc(job, data, index), str(path))
            doc = Document(path)
        p = next(p for p in doc.paragraphs if p.text == job["summary_text"])
        self.assertTrue(all(r.bold and r.italic and r.underline for r in p.runs))
        self.assertEqual(p.runs[0].font.size.pt, 12)
        self.assertEqual(p.paragraph_format.space_before.pt, 10)
        self.assertEqual(p.paragraph_format.space_after.pt, 8)
        self.assertAlmostEqual(p.paragraph_format.line_spacing, 1.2)
        self.assertTrue(p.paragraph_format.keep_with_next)
        self.assertEqual(p.paragraph_format.left_indent, proposal.DESC_TAB)
        self.assertAlmostEqual(p.paragraph_format.first_line_indent.inches,
                               proposal.QTY_TAB.inches - proposal.DESC_TAB.inches, places=3)

    def test_bad_format_is_rejected(self):
        for value in ({"font_size": 0}, {"line_spacing": float('nan')}, {"start": "bad"}, {"page_break": "yes"}, []):
            with self.subTest(value=value), self.assertRaises(ValueError):
                quote_format.validate(value)


if __name__ == "__main__":
    unittest.main()
